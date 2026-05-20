import io
import sys
import time
import json
import copy
import shutil
import signal
import logging
import datetime
import zipfile
import tempfile
import openpyxl
import openpyxl.utils
import openpyxl.styles
import threading
import subprocess as sp
from art import text2art
from pathlib import Path
from collections import deque
from itertools import zip_longest
from typing import Tuple, List, Dict, Any
from types import MappingProxyType

__IS_ALIVE__ = True

__SPECIAL_CMD__ = MappingProxyType({
    "results": ["python3", "host_ssh_script.py", "--results"],
    "brownout": ["python3", "host_ssh_script.py", "--test", "brownout", "--serial_test"]
})


def signal_handler(_sig, _frame):
    global __IS_ALIVE__
    __IS_ALIVE__ = False


signal.signal(signal.SIGINT, signal_handler)
signal.signal(signal.SIGTERM, signal_handler)


class BlockingQueue:
    def __init__(self, maxsize=0):
        self.maxsize = maxsize
        self.queue = deque(maxlen=maxsize) if maxsize > 0 else deque()
        self.lock = threading.Lock()
        self.not_empty = threading.Condition(self.lock)

    def put(self, item):
        with self.lock:
            if 0 < self.maxsize <= len(self.queue):
                self.queue.pop()
            self.queue.appendleft(item)
            self.not_empty.notify()

    def get(self) -> Any:
        with self.lock:
            while not self.queue:
                self.not_empty.wait()
            return self.queue.popleft()


__INFO_EXCEL_PATH__ = Path().home().joinpath("Desktop/FactoryData/.script_test_info.xlsx")
__NECESSARY_FILE__ = ("common_files", "project_files", "test_flow.ini", "config.json", "host_ssh_script.py")
info_queue = BlockingQueue(maxsize=64)


def scan_local_script() -> Dict:
    script_path_dict = {}
    desktop = Path().home().joinpath("Desktop")
    patterns = (
        '*_*_Factory',
        'DOE'
    )
    found_scripts = []
    for pattern in patterns:
        for dir_path in desktop.glob(pattern):
            if dir_path.is_dir():
                for script_file in dir_path.iterdir():
                    if script_file.is_dir() and all([
                        necessary_file in
                        [file.name for file in script_file.iterdir()]
                        for necessary_file in __NECESSARY_FILE__
                    ]):
                        found_scripts.append(script_file)
    for script_path in found_scripts:
        if script_path.name not in script_path_dict.keys():
            script_path_dict[script_path.name] = [str(script_path)]
        else:
            script_path_dict[script_path.name].append(str(script_path))
    return script_path_dict


def __get_cmd__(test_method: str, ecid: str) -> List[str]:
    test_method_split = test_method.split()
    __COMMON_CMD__ = ["python3", "host_ssh_script.py", "--test", test_method_split]
    return __SPECIAL_CMD__.get(test_method, __COMMON_CMD__) + ["--ecid", ecid]


def log_message(base_path: Path, test_method: str, message: str):
    log_filename = test_method + '_' + datetime.datetime.now().strftime("%Y-%m-%dT%H%M%S") + ".log"

    logger = logging.getLogger("console_logger")
    logger.setLevel(logging.DEBUG)

    file_handler = logging.FileHandler(base_path.joinpath("logs", log_filename))
    formatter = logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.info(message)

    def archive_and_cleanup_old_logs():
        logs_dir = base_path.joinpath("logs")
        archive_dir = base_path.joinpath("logs", "archive")
        archive_dir.mkdir(parents=True, exist_ok=True)
        three_days_ago = datetime.datetime.now() - datetime.timedelta(days=3)
        log_by_date = {}
        for log_file in logs_dir.glob("*.log"):
            if log_file.name == "archive":
                continue
            file_mtime = datetime.datetime.fromtimestamp(log_file.stat().st_mtime)
            if file_mtime < three_days_ago:
                date_str = file_mtime.strftime("%Y-%m-%d")
                if date_str not in log_by_date:
                    log_by_date[date_str] = []
                log_by_date[date_str].append(log_file)
        for date_str, files in log_by_date.items():
            zip_filename = f"{date_str}.zip"
            zip_path = archive_dir.joinpath(zip_filename)
            with zipfile.ZipFile(zip_path, "w") as zipf:
                for file in files:
                    zipf.write(file, arcname=file.name)
            for file in files:
                file.unlink(missing_ok=True)
        three_months_ago = datetime.datetime.now() - datetime.timedelta(days=90)
        for log_file in logs_dir.glob("*.log"):
            file_mtime = datetime.datetime.fromtimestamp(log_file.stat().st_mtime)
            if file_mtime < three_months_ago:
                log_file.unlink(missing_ok=True)
        for archive_file in archive_dir.glob("*.zip"):
            file_mtime = datetime.datetime.fromtimestamp(archive_file.stat().st_mtime)
            if file_mtime < three_months_ago:
                archive_file.unlink(missing_ok=True)
    archive_and_cleanup_old_logs()


def save_info_to_excel():
    if __INFO_EXCEL_PATH__.exists():
        wb = openpyxl.load_workbook(__INFO_EXCEL_PATH__)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'sync_info'
        title = ('serial_number', 'test_name', 'unit_number', 'config', 'bundle', 'script_path', 'battery',
                 'sync_time')
        ws.append(title)
        header_font = openpyxl.styles.Font(name='Calibri', size=16, bold=True)
        header_fill = openpyxl.styles.PatternFill(start_color='9BC2E6', end_color='9BC2E6', fill_type='solid')
        thin = openpyxl.styles.Side(style='thin', color='000000')
        border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = len(cell.value) * 1.5
        wb.save(__INFO_EXCEL_PATH__)
    data_row = 2
    thin_side = openpyxl.styles.Side(style='thin', color='000000')
    data_font = openpyxl.styles.Font(size=16)
    data_fill = openpyxl.styles.PatternFill(start_color='EEEEEE', end_color='EEEEEE', fill_type='solid')
    while __IS_ALIVE__:
        info = info_queue.get()
        for ecid, data in info.items():
            sn = data.get("serial_number", "")
            test_name = data.get("test_method", "")
            unit_number = data.get("unit_number", "")
            config = data.get("config", "")
            sw_version = data.get("sw_vers", "")
            script_path = data.get("script_path", "")
            battery = data.get("battery", "")
            sync_time = data.get("sync_time", "")
            info_data = [sn, test_name, unit_number, config, sw_version, script_path, battery, sync_time]
            ws.insert_rows(data_row)
            for col, value in enumerate(info_data, start=1):
                cell = ws.cell(row=data_row, column=col, value=value)
                cell.font = data_font
                cell.fill = data_fill
                cell.border = openpyxl.styles.Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            for col_index, value in enumerate(info_data, start=1):
                col_letter = openpyxl.utils.get_column_letter(col_index)
                current_col_dim = ws.column_dimensions.get(col_letter)
                current_width = current_col_dim.width if current_col_dim else None
                required_width = len(str(value)) * 1.5
                if current_width is None or current_width < required_width:
                    ws.column_dimensions[col_letter].width = required_width
            data_row += 1
        wb.save(__INFO_EXCEL_PATH__)
    wb.close()


def is_json_file_valid(file_path: str | Path) -> bool:
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            json.load(file)
        return True
    except json.JSONDecodeError as e:
        print(f"JSON format error: {e}")
        return False
    except FileNotFoundError:
        print("cannot find json file")
        return False
    except Exception as e:
        print(f"Unknown error: {e}")
        return False


class TestScript:
    def __init__(self, base_path: Path):
        self.__base_path__ = base_path
        self.__script_path__ = {}
        self.__testing_dict__ = {}
        self.__testing_data_size__ = {}
        self.__running__ = True
        self.__load_config__()
        self.__check_config_health_thread__ = threading.Thread(target=self.__check_config_health__, daemon=True)
        self.__test_lifetime_thread__ = threading.Thread(target=self.__test_lifetime_detector__, daemon=True)
        self.__info_record_thread__ = threading.Thread(target=save_info_to_excel)
        self.__check_config_health_thread__.start()
        self.__test_lifetime_thread__.start()
        self.__info_record_thread__.start()
        self.__thread_join_callable__ = [
            self.__check_config_health_thread__.join,
            self.__test_lifetime_thread__.join,
            self.__info_record_thread__.join
        ]

    def stop(self):
        self.__running__ = False
        for join in self.__thread_join_callable__:
            join()

    def __check_config_health__(self):
        while self.__running__:
            __script_change__ = False
            __removable_scripts__ = []
            for script_name, data in self.__script_path__.items():
                if Path(data["path"]).exists() and not data["status"]:
                    self.__script_path__[script_name]["status"] = True
                    __script_change__ = True
                elif not Path(data["path"]).exists() and data["status"]:
                    self.__script_path__[script_name]["status"] = False
                    __script_change__ = True
                    __removable_scripts__.append(script_name)
            for removable_script in __removable_scripts__:
                self.__terminate_by_script_name__(removable_script)
            if __script_change__:
                self.__save_config__()
            time.sleep(0.75)

    def __terminate_by_script_name__(self, script_name: str):
        __removable_ecids__ = []
        for ecid, data in self.__testing_dict__.items():
            if data.get("test_status", {}).get("script_name", None) == script_name:
                process = data.get("process", None)
                if isinstance(process, sp.Popen):
                    process.terminate()
                    __removable_ecids__.append(ecid)
        for removable_ecid in __removable_ecids__:
            self.__testing_dict__.pop(removable_ecid, None)

    def __save_config__(self):
        config_path = Path(self.__base_path__).joinpath("config/test_scripts.json")
        with open(config_path, "w") as f:
            json.dump(self.__script_path__, f)

    def __load_config__(self):
        config_path = Path(self.__base_path__).joinpath("config/test_scripts.json")
        if config_path.exists():
            with open(config_path, "r") as f:
                self.__script_path__ = json.load(f)

    def __get_test_methods__(self, script_name: str = "", script_path: str = "") -> Tuple[bool, List]:
        if script_name:
            script_base = self.__script_path__.get(script_name, {}).get("path", None)
        elif script_path:
            script_base = Path(script_path)
        else:
            return False, []
        if not Path(script_base).exists() or not Path(script_base).is_dir():
            return False, []
        methods_config = script_base.joinpath("methods.json")
        test_app_path = script_base.joinpath("TestApp")
        if methods_config.exists():
            with open(methods_config, "r") as f:
                methods = json.load(f)
                return True, methods
        elif test_app_path.exists():
            files = [f.name for f in test_app_path.iterdir()]
            files.remove("收集数据")
            files = [s.lower() for s in files]
            files.append("results")
            files.sort()
            return True, files
        else:
            return True, []

    def add_script_path(self, script_name: str, script_path: str) -> Tuple[bool, str]:
        if not Path(script_path).exists() or not Path(script_path).is_dir():
            desktop = Path().home().joinpath("Desktop")
            patterns = [
                '*_*_Factory',
                'DOE'
            ]
            found_scripts = []
            for pattern in patterns:
                for dir_path in desktop.glob(pattern):
                    if dir_path.is_dir():
                        script = dir_path.joinpath(script_path)
                        if script.exists() and script.is_dir():
                            found_scripts.append(script)
            if len(found_scripts) == 0:
                return False, "Script path dose not exists"
            elif len(found_scripts) > 1:
                return False, "Multiple scripts found, please input full path"
            else:
                script_path = str(found_scripts[0])
        file_list = [file.name for file in Path(script_path).iterdir()]
        for necessary_file in __NECESSARY_FILE__:
            if necessary_file not in file_list:
                return False, "Script unlaw"
        if script_name not in self.__script_path__.keys():
            result, methods = self.__get_test_methods__(script_path=script_path)
            if result:
                self.__script_path__[script_name] = {"path": script_path, "status": True, "methods": methods}
                self.__save_config__()
            else:
                return False, "Value error"
            return True, "OK"
        else:
            return False, "Script name already exists"

    def remove_script_path(self, script_name: str = "", script_path: str = "") -> Tuple[bool, str]:
        if script_name:
            result = self.__script_path__.pop(script_name, None)
        elif script_path:
            self.__script_path__ = {k: v for k, v in self.__script_path__.items() if v["path"] != script_path}
            result = True, "OK"
        else:
            result = False, "Unexpected error"
        if result:
            self.__save_config__()
            return True, "OK"
        else:
            return False, "Unexpected error"

    def get_script_info(self) -> Dict:
        info = copy.deepcopy(self.__script_path__)
        for script_name, data in info.items():
            data.pop("methods", None)
        return info

    def get_script_methods(self, script_name: str = "", script_path: str = "") -> Tuple[bool, Dict]:
        methods_dict = {}
        if script_name:
            methods = self.__script_path__.get(script_name, {}).get("methods", None)
            if methods is None:
                return False, {}
            methods_dict[script_name] = methods
        elif script_path:
            for script_name, data in self.__script_path__.items():
                if data.get("path", None) == script_path:
                    methods = data.get("methods", None)
                    if methods:
                        methods_dict[script_name] = methods
        else:
            return False, {}
        if methods_dict:
            return True, methods_dict
        else:
            return False, {}

    def get_testing_ecids(self):
        testing_ecids = []
        for ecid, data in self.__testing_dict__.items():
            thread = data.get("thread", None)
            if isinstance(thread, threading.Thread):
                if thread.is_alive():
                    testing_ecids.append(ecid)
        return testing_ecids

    def get_log_stream(self, ecids: List[str]):
        log_buffers = []
        threads = []
        buffer_positions = []
        for ecid in ecids:
            data = self.__testing_dict__.get(ecid)
            if data:
                log_buffers.append(data["log_stream"])
                threads.append(data["thread"])
                buffer_positions.append(0)
            else:
                print(f"ECID {ecid} not found")
        if not log_buffers:
            return
        stream_contents = []
        while any([thread.is_alive() for thread in threads]) and self.__running__:
            stream_contents.clear()
            for i, buffer in enumerate(log_buffers):
                if isinstance(buffer, io.BytesIO):
                    if not buffer.closed:
                        content = buffer.getvalue().decode('utf-8')
                    else:
                        continue
                elif isinstance(buffer, io.StringIO):
                    if not buffer.closed:
                        content = buffer.getvalue()
                    else:
                        continue
                else:
                    raise ValueError(f"Unsupported buffer type: {type(buffer)}")
                new_content = content[buffer_positions[i]:]
                buffer_positions[i] = len(content)
                stream_contents.append(new_content.splitlines())
            for lines in zip_longest(*stream_contents):
                line_str = "\n"
                for line in lines:
                    if line:
                        if line.strip():
                            line_str = line_str + line.strip() + '\n'
                yield line_str.encode('utf-8')
            has_content = any(any(lines) for lines in stream_contents)
            if not has_content:
                yield ""
            time.sleep(0.1)
        yield "################################################\n"
        yield text2art("Complete")

    def run_test_script(self, test_method: str, script_name: str, units: dict) -> Tuple[bool, str]:
        if not self.__script_path__:
            return False, "Script path empty"
        if test_method == "results":
            self.__testing_data_size__.update({ecid: data.get("data_size", "0K") for ecid, data in units.items()})
            if not self.__adjust_space__():
                return False, "Space not enough"
        order = -1
        for ecid, data in units.items():
            order += 1
            info = {ecid: data}
            info[ecid]["test_method"] = test_method
            info[ecid]["script_path"] = self.__script_path__.get(script_name, {}).get("path", "UNDEFINED")
            log_buffer = io.BytesIO()
            thread = threading.Thread(target=self.__run_process__,
                                      args=(test_method, ecid, log_buffer, script_name, info,))
            self.__testing_dict__[ecid] = {
                "workdir": tempfile.mkdtemp(),
                "order": order,
                "log_stream": log_buffer,
                "thread": thread,
                "process": None,
                "test_status": {
                    "log_level": "info",
                    "script_name": script_name,
                    "test_method": test_method
                }
            }
            thread.start()
        return True, "OK"

    def terminate_script(self, units: list[str]) -> Tuple[bool, list[str]]:
        success_ecid = []
        for ecid in units:
            process = self.__testing_dict__.get(ecid, {}).get("process", None)
            if isinstance(process, sp.Popen):
                process.terminate()
                success_ecid.append(ecid)
        if len(success_ecid) == len(units):
            return True, success_ecid
        else:
            return False, success_ecid

    def __adjust_space__(self) -> bool:
        retain_scale = 0.5
        retain_space = 2 * 1024
        total, used, free = shutil.disk_usage("/")
        host_size_str = f"{free // (1024**3)}G"
        if host_size_str.endswith("T"):
            host_size = float(host_size_str[:-1]) * 1024**2
        elif host_size_str.endswith("G"):
            host_size = float(host_size_str[:-1]) * 1024
        elif host_size_str.endswith("M"):
            host_size = float(host_size_str[:-1])
        else:
            return False
        if host_size == 0:
            return True
        all_data_size = 0
        for _, data_size_str in self.__testing_data_size__.items():
            if data_size_str.endswith("T"):
                data_size = float(data_size_str[:-1]) * 1024**2
            elif data_size_str.endswith("G"):
                data_size = float(data_size_str[:-1]) * 1024
            elif data_size_str.endswith("M"):
                data_size = float(data_size_str[:-1])
            else:
                data_size = 0
            all_data_size += data_size
        if any((host_size - all_data_size) < val for val in [all_data_size * retain_scale, retain_space]):
            return False
        else:
            return True

    @staticmethod
    def __pre_process_test_method__(test_method: str) -> Tuple[bool, str]:
        if '_' not in test_method:
            return True, test_method
        test_method_split = test_method.split('_')
        if "arcas" in test_method_split[-1]:
            test_method = ' '.join(list(reversed(test_method_split)))
        else:
            return False, "No Matched"
        return True, test_method

    def __run_process__(self, test_method: str, ecid: str, log_buffer: io.BytesIO, script_name: str,
                        info: dict) -> bool:
        if not self.__script_path__:
            return False
        if not self.__script_path__[script_name]["status"]:
            return False
        script_dir = self.__script_path__[script_name]["path"]
        methods = self.__script_path__[script_name]["methods"]
        script_name_str = text2art(script_name)
        workdir = self.__testing_dict__.get(ecid, {}).get("workdir", script_dir)
        order = self.__testing_dict__.get(ecid, {}).get("order", 0)
        log_buffer.write(script_name_str.encode("utf-8"))
        log_buffer.write(f"Script path: {script_dir}\n".encode("utf-8"))
        log_buffer.write(f"Workdir: {workdir}\n".encode("utf-8"))
        log_buffer.write(b"================")
        log_buffer.write(b"================")
        log_buffer.write(b"================")
        log_buffer.write(b"================")
        log_buffer.write(b"\n")
        if methods and test_method not in methods:
            return False
        process_ok, data = self.__pre_process_test_method__(test_method)
        if not process_ok:
            return False
        test_method = data
        cmd = __get_cmd__(test_method, ecid)
        if not cmd:
            return False
        if str(workdir) != str(script_dir) and test_method not in __SPECIAL_CMD__.keys():
            shutil.copytree(
                script_dir,
                workdir,
                dirs_exist_ok=True,
                ignore=shutil.ignore_patterns('.DS_Store', 'TestApp', 'README.md')
            )
        else:
            workdir = script_dir
        time.sleep(abs(order * 0.15))
        process = sp.Popen(
            cmd,
            stdout=sp.PIPE,
            stderr=sp.PIPE,
            cwd=workdir,
        )
        self.__testing_dict__[ecid]["process"] = process
        for line in process.stdout:
            log_buffer.write(line)
            if not self.__running__:
                break
        if self.__running__:
            process.wait()
        else:
            process.terminate()
        if process.returncode == 0:
            info[ecid]["sync_time"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            info_queue.put(info)
        return True

    def __test_lifetime_detector__(self):
        while self.__running__:
            __removable_ecids__ = []
            for ecid, data in self.__testing_dict__.items():
                thread = data.get("thread", None)
                log_level = data.get("test_status", {}).get("log_level", None)
                test_method = data.get("test_status", {}).get("test_method", None)
                log_stream = data.get("log_stream", None)
                if all([isinstance(thread, threading.Thread),
                        isinstance(log_level, str),
                        isinstance(test_method, str),
                        isinstance(log_stream, io.BytesIO) or isinstance(log_stream, io.StringIO)]):
                    if not thread.is_alive():
                        __removable_ecids__.append(ecid)
                else:
                    __removable_ecids__.append(ecid)
            for ecid in self.__testing_data_size__.keys():
                data = self.__testing_dict__.get(ecid, {})
                thread = data.get("thread", None)
                if isinstance(thread, threading.Thread):
                    if not thread.is_alive():
                        __removable_ecids__.append(ecid)
                else:
                    __removable_ecids__.append(ecid)
            __removable_ecids__ = list(set(__removable_ecids__))
            for removable_ecid in __removable_ecids__:
                log_level = self.__testing_dict__.get(removable_ecid, {}).get("test_status", {}).get("log_level", None)
                log_stream = self.__testing_dict__.get(removable_ecid, {}).get("log_stream", None)
                test_method = self.__testing_dict__.get(removable_ecid, {}).get("test_status", {}).get("test_method", None)
                workdir = self.__testing_dict__.get(removable_ecid, {}).get("workdir", None)
                if Path(workdir).is_dir():
                    shutil.rmtree(workdir)
                if log_level == "info":
                    log_str = log_stream.getvalue().decode('utf-8', errors='ignore')
                    log_message(self.__base_path__, test_method, log_str)
                log_stream.close()
                self.__testing_dict__.pop(removable_ecid, None)
                self.__testing_data_size__.pop(removable_ecid, None)
            time.sleep(0.75)


if __name__ == '__main__':
    if getattr(sys, 'frozen', False):
        __base_path__ = Path(sys.executable).parent
    else:
        __base_path__ = Path(__file__).parent
    test_script = TestScript(__base_path__)
    script_str = input("Input script name and script path(format script_name:script_path)> ").strip()
    add_script_results = True
    if script_str:
        add_script_results = test_script.add_script_path(script_name=script_str.split(":")[0],
                                                         script_path=script_str.split(":")[1])
    if not add_script_results:
        print("Script path not found")
        sys.exit(0)
    run_str = input("Input run script param(format test_method:script_name:['units'])> ").strip()
    run_units = json.loads(run_str.split(":")[2])
    run_script_result = test_script.run_test_script(test_method=run_str.split(":")[0],
                                                    script_name=run_str.split(":")[1], units=run_units)
    if not run_script_result:
        print("Script run failed")
        sys.exit(0)
    print(test_script.get_log_stream(ecids=run_units))
    test_script.stop()
    sys.exit(0)
